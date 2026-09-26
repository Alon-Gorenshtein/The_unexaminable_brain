"""Guards for the revision packet. Every check returns a list of problems (empty means pass).

    python3 code/tools/check_packet.py PACKET_DIR

PACKET_DIR holds 1_Cover_Letter.docx, 2_Response_to_Reviewers.docx, 3_Manuscript_COLORED_CHANGES.docx,
3b_Manuscript_CLEAN.docx, 4_Supplement.docx, 5_Reporting_Checklist_STROBE_RECORD.docx and Figures/*.tif. Prints "packet checks pass" and exits 0
when every guard passes; otherwise prints one line per problem and exits 1. Either way it then lists
"open item:" lines for text the authors must still supply; those never change the exit code.
"""
import csv
import json
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "code"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

COVER, RESPONSE = "1_Cover_Letter.docx", "2_Response_to_Reviewers.docx"
COLORED, CLEAN, SUPPLEMENT = "3_Manuscript_COLORED_CHANGES.docx", "3b_Manuscript_CLEAN.docx", "4_Supplement.docx"
CHECKLIST = "5_Reporting_Checklist_STROBE_RECORD.docx"
DOCS = (COVER, RESPONSE, COLORED, CLEAN, SUPPLEMENT, CHECKLIST)
SOURCES = {COLORED: "manuscript.md", CLEAN: "manuscript.md", SUPPLEMENT: "supplement.md",
           RESPONSE: "response_to_reviewers.md", COVER: "cover_letter_revision.md",
           CHECKLIST: "reporting_checklist.md"}
CLAIM_FILES = {"manuscript/manuscript.md": CLEAN, "manuscript/supplement.md": SUPPLEMENT,
               "manuscript/response_to_reviewers.md": RESPONSE, "manuscript/cover_letter_revision.md": COVER}
# Wording the revision removed. The response quotes the reviewer and the submitted text, so there
# these may appear inside block quotes only.
REMOVED_TERMS = ["misclassif", "general property", "true GCS", "legacy"]
TERMS_FILE = ROOT / "docs" / "revision" / "packet_terms.tsv"
SIGNATURE = "[signature"
DECLARATION_HEADINGS = ["Acknowledgments", "Author Contributions", "Statements and Declarations",
                        "Ethical considerations", "Consent to participate", "Consent for publication",
                        "Declaration of conflicting interest", "Funding statement", "Data availability"]
# Text the authors must still supply; listed after a passing check, never a failure.
HUMAN_MARKERS = ["[CRediT roles to be added by the authors]"]


def load_terms(path=TERMS_FILE):
    """{"process": [...], "metadata": [...]} from the tab-separated terms file (kind, term)."""
    terms = {"process": [], "metadata": []}
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        if line.strip() and not line.startswith("#"):
            kind, term = line.split("\t", 1)
            terms[kind].append(term)
    return terms


def _xml(docx, name):
    with zipfile.ZipFile(docx) as z:
        return z.read(name).decode("utf8") if name in z.namelist() else ""


def _plain(docx):
    return subprocess.run(["pandoc", "-f", "docx", "-t", "plain", "--wrap=none", str(docx)],
                          capture_output=True, text=True, check=True).stdout


def _body_paragraphs(docx):
    """(style name, text, inside a table) for every paragraph of the body, tables included."""
    from docx import Document
    doc = Document(docx)
    out = []

    def walk(container, in_table):
        for p in container.paragraphs:
            out.append((p.style.name if p.style is not None else "", p.text, in_table))
        for t in container.tables:
            for row in t.rows:
                for cell in row.cells:
                    walk(cell, True)
    walk(doc, False)
    return out


def has_page_field(docx):
    with zipfile.ZipFile(docx) as z:
        foot = "".join(z.read(n).decode("utf8") for n in z.namelist() if n.startswith("word/footer"))
    ok = re.search(r"instrText[^>]*>\s*PAGE\b|fldSimple[^>]*PAGE", foot)
    return [] if ok else [f"{Path(docx).name}: no PAGE field in any footer"]


def _paragraph_lines(docx):
    """w:line value set directly on each body paragraph outside tables (None when not set)."""
    from lxml import etree
    root = etree.fromstring(_xml(docx, "word/document.xml").encode("utf8"))
    body = root.find(f"{W_NS}body")
    vals = []
    for p in body.iter(f"{W_NS}p"):
        if any(a.tag == f"{W_NS}tbl" for a in p.iterancestors()):
            continue
        sp = p.find(f"{W_NS}pPr/{W_NS}spacing")
        vals.append(sp.get(f"{W_NS}line") if sp is not None else None)
    return vals


def is_double_spaced(docx):
    """Double spacing from the document defaults or the body, and no body paragraph overriding it."""
    dd = re.search(r"<w:docDefaults>.*?</w:docDefaults>", _xml(docx, "word/styles.xml"), re.S)
    ok = 'w:line="480"' in _xml(docx, "word/document.xml") or bool(dd and 'w:line="480"' in dd.group(0))
    problems = [] if ok else [f"{Path(docx).name}: no double spacing (w:line=480) in the document defaults or body"]
    other = sorted({v for v in _paragraph_lines(docx) if v not in (None, "480")})
    if other:
        problems.append(f"{Path(docx).name}: body paragraphs override double spacing (w:line {other})")
    return problems


def is_single_spaced(docx):
    vals = _paragraph_lines(docx)
    bad = sum(1 for v in vals if v != "240")
    return [f"{Path(docx).name}: {bad} of {len(vals)} body paragraphs are not single-spaced"] if bad or not vals else []


def has_no_embedded_figures(docx):
    with zipfile.ZipFile(docx) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    return [f"{Path(docx).name}: embedded media: {media}"] if media else []


def colored_run_count(docx, color="C00000"):
    return len(re.findall(rf'<w:color w:val="{color}"', _xml(docx, "word/document.xml")))


def highlight_count(docx):
    return len(re.findall(r"<w:highlight ", _xml(docx, "word/document.xml")))


def title_matches(docx, title_file):
    title = Path(title_file).read_text().strip()
    return [] if title in _plain(docx).replace("\n", " ") else [f"title not found in {Path(docx).name}"]


def not_stale(docx, md):
    return [] if Path(docx).stat().st_mtime >= Path(md).stat().st_mtime else [f"{docx} is older than {md}"]


def strings_absent(docx, strings, ignore_case=False):
    t = _plain(docx).replace("\n", " ")
    if ignore_case:
        return [f"{Path(docx).name}: forbidden string present: {s!r}" for s in strings if s.lower() in t.lower()]
    return [f"{Path(docx).name}: forbidden string present: {s!r}" for s in strings if s in t]


def strings_present(docx, strings):
    t = _plain(docx).replace("\n", " ")
    return [f"{Path(docx).name}: required string missing: {s!r}" for s in strings if s not in t]


def strings_only_in_quotes(docx, strings):
    """The strings may appear only in block quotes (pandoc's "Block Text" style)."""
    bad = []
    for style, text, _ in _body_paragraphs(docx):
        if style == "Block Text":
            continue
        low = text.lower()
        bad += [f"{Path(docx).name}: {s!r} outside a block quote: {text[:80]!r}" for s in strings if s.lower() in low]
    return bad


def signature_placeholder_once(docx, expected):
    """The signature placeholder appears exactly `expected` times (0 in every document: none is left)."""
    n = _plain(docx).lower().count(SIGNATURE)
    return [] if n == expected else [f"{Path(docx).name}: signature placeholder appears {n} times, expected {expected}"]


def abstract_words(docx):
    """Words of the abstract, section labels included, between the Abstract heading and the Keywords line
    (-1 when either is missing, or the Keywords line comes first)."""
    text = _plain(docx)
    a = re.search(r"^\s*Abstract\s*$", text, re.M)
    k = re.compile(r"^\s*Keywords\s*:", re.M).search(text, a.end()) if a else None
    if not a or not k:
        return -1
    return len(text[a.end():k.start()].split())


def abstract_within_limit(docx, lo=250, hi=300):
    n = abstract_words(docx)
    if n < 0:
        return [f"{Path(docx).name}: abstract or keywords line not found"]
    return [] if lo <= n <= hi else [f"{Path(docx).name}: abstract has {n} words; the journal requires {lo}-{hi}"]


def declarations_present(docx):
    """Each declaration heading stands on a line of its own."""
    lines = [l.strip() for l in _plain(docx).splitlines()]
    return [f"{Path(docx).name}: missing heading {h!r}" for h in DECLARATION_HEADINGS if h not in lines]


def open_human_items(root):
    """Markers of text the authors must still supply, as "<file>: <marker>"."""
    found = []
    for n in (COLORED, CLEAN, COVER, RESPONSE, SUPPLEMENT):
        p = Path(root) / n
        if p.is_file():
            found += [f"{n}: {m}" for m in HUMAN_MARKERS if m in _plain(p)]
    return found


def frontmatter_wordcount_matches(docx):
    """The "Word count: N" line of the front matter equals the strict count of the built file."""
    from strict_wordcount import strict_count
    m = re.search(r"Word count:\s*([\d,]+)", _plain(docx))
    if not m:
        return [f"{Path(docx).name}: no 'Word count:' line in the front matter"]
    stated, measured = int(m.group(1).replace(",", "")), strict_count(str(docx))
    return [] if stated == measured else [f"{Path(docx).name}: front matter says {stated} words, strict count is {measured}"]


def no_forbidden_metadata(docx, patterns=()):
    """Document properties carry no email address, no "python-docx" and nothing matching `patterns`."""
    rx = re.compile("|".join([r"[\w.+-]+@[\w-]+\.[\w.]+", "python-docx", *patterns]))
    bad = []
    for part in ("docProps/core.xml", "docProps/app.xml", "docProps/custom.xml"):
        text = re.sub(r'xmlns(:\w+)?="[^"]*"', "", _xml(docx, part))  # namespace URLs are not metadata
        for m in rx.finditer(text):
            bad.append(f"{Path(docx).name}: {part} contains {m.group(0)!r}")
    return bad


def no_sidecar_files(root):
    """No AppleDouble (._*) or .DS_Store files in the packet, and none inside any .docx archive."""
    root = Path(root)
    bad = [f"sidecar file: {p.relative_to(root)}" for p in root.rglob("*")
           if p.name.startswith("._") or p.name == ".DS_Store"]
    for d in root.rglob("*.docx"):
        if d.name.startswith("._"):
            continue
        with zipfile.ZipFile(d) as z:
            bad += [f"{d.name}: archive entry {n}" for n in z.namelist() if Path(n).name.startswith("._")]
    return bad


def no_track_changes(docx):
    bad = []
    with zipfile.ZipFile(docx) as z:
        for n in z.namelist():
            if not (n.startswith("word/") and n.endswith(".xml")):
                continue
            x = z.read(n).decode("utf8", "ignore")
            for pat in (r"<w:ins\b", r"<w:del\b", r"<w:moveFrom\b", r"<w:moveTo\b", r"<w:comment\b",
                        r"<w:commentReference\b", r"<w:commentRangeStart\b"):
                if re.search(pat, x):
                    bad.append(f"{Path(docx).name}: {n} contains {pat[1:-2]}")
    return bad


def tiff_ok(fig_dir, png_dir, mapping, dpi=600):
    """Each journal-named TIFF exists, is 600-dpi RGB without alpha, and has the pixels of its source PNG."""
    from PIL import Image, ImageChops
    bad = []
    for stem, name in mapping.items():
        tif, png = Path(fig_dir) / f"{name}.tif", Path(png_dir) / f"{stem}.png"
        if not tif.is_file():
            bad.append(f"missing figure {tif.name}")
            continue
        with Image.open(tif) as t, Image.open(png) as s:
            got = t.info.get("dpi", (0, 0))
            if [round(float(v)) for v in got] != [dpi, dpi]:
                bad.append(f"{tif.name}: {got} dpi, expected {dpi}")
            if t.mode != "RGB":
                bad.append(f"{tif.name}: mode {t.mode}, expected RGB without alpha")
            if t.size != s.size:
                bad.append(f"{tif.name}: {t.size} px, source {png.name} is {s.size}")
            elif ImageChops.difference(t.convert("RGB"), s.convert("RGB")).getbbox() is not None:
                bad.append(f"{tif.name}: pixels differ from {png.name} (stale export)")
    return bad


def supplement_numbering(docx, n_tables=18, n_figures=3):
    """eTable 1..n and eFigure 1..m captions (paragraphs that start with their number), in order."""
    got = []
    for style, text, in_table in _body_paragraphs(docx):
        m = re.match(r"\s*(eTable|eFigure) (\d+)\.", text)
        if m and not in_table:
            got.append(f"{m.group(1)} {m.group(2)}")
    want = [f"eTable {i}" for i in range(1, n_tables + 1)] + [f"eFigure {i}" for i in range(1, n_figures + 1)]
    return [] if got == want else [f"{Path(docx).name}: captions {got} do not match {want[0]}..{want[-1]} in order"]


def docx_claim_text(docx):
    """Text of a .docx for the claims registry: every paragraph's text, and every table row written as
    a markdown pipe row ("| a | b |"), whitespace collapsed. Superscripts are plain characters here."""
    from docx import Document
    doc = Document(docx)
    parts = []

    def walk(container):
        for block in container.iter_inner_content():
            if hasattr(block, "rows"):
                for row in block.rows:
                    parts.append("| " + " | ".join(c.text for c in row.cells) + " |")
            else:
                parts.append(block.text)
    walk(doc)
    return re.sub(r"\s+", " ", " ".join(parts))


def docx_literal(lit):
    """The registered literal as it reads in a .docx: markdown emphasis markers and the carets of
    superscripts removed (their content kept), whitespace collapsed."""
    t = re.sub(r"\^([^^\s]+)\^", r"\1", lit)
    t = t.replace("**", "").replace("`", "")
    t = re.sub(r"(?<![\w*])\*(?=\S)|(?<=\S)\*(?![\w*])", "", t)
    return re.sub(r"\s+", " ", t).strip()


def claims_in_docx(packet, claims=ROOT / "output" / "revision" / "claims.tsv", scratch=None):
    """Run the claims registry against the built .docx files instead of the markdown."""
    from check_claims import check_rows, parse_claims
    from lib_digest import build_digest, stale_stems
    packet = Path(packet)
    with open(claims, newline="", encoding="utf-8") as fh:
        rows, bad = parse_claims(fh)
    mapped = []
    for f, key, kind, lit in rows:
        if f not in CLAIM_FILES:
            bad.append(f"{key}: no built document for {f}")
            continue
        mapped.append((str(packet / CLAIM_FILES[f]), key, kind, docx_literal(lit)))
    out = Path(scratch or tempfile.mkdtemp()) / "claims_docx.tsv"
    with open(out, "w", newline="", encoding="utf-8") as fh:
        csv.writer(fh, delimiter="\t", quoting=csv.QUOTE_NONE, escapechar="\\").writerows(mapped)
    digest = json.loads((ROOT / "output" / "revision" / "revision_digest.json").read_text(encoding="utf-8"))
    stale = stale_stems(build_digest(ROOT / "output" / "revision"), digest)
    if stale:
        return bad + ["digest is stale (re-run code/72_revision_digest.py): " + ", ".join(stale)]
    cache = {}

    def read(f):
        if f not in cache:
            cache[f] = docx_claim_text(f)
        return cache[f]
    return bad + check_rows(mapped, digest, read)


def check(root):
    from export_figures import FIG, MAP
    root = Path(root)
    md = ROOT / "manuscript"
    title = md / "title.txt"
    docs = {n: root / n for n in DOCS}
    missing = [n for n, p in docs.items() if not p.is_file()]
    if missing:
        return [f"missing {n}" for n in missing]
    ms, clean = docs[COLORED], docs[CLEAN]
    terms = load_terms()
    problems = []
    for n in DOCS:
        problems += has_page_field(docs[n])
    for n in (COLORED, CLEAN, SUPPLEMENT):
        problems += is_double_spaced(docs[n])
    for n in (COVER, RESPONSE, CHECKLIST):  # the lab's earlier response letter is single-spaced; a letter is too
        problems += is_single_spaced(docs[n])
    for n in DOCS:
        problems += title_matches(docs[n], title) + not_stale(docs[n], md / SOURCES[n])
        problems += no_forbidden_metadata(docs[n], terms["metadata"]) + no_track_changes(docs[n])
        problems += strings_absent(docs[n], terms["process"])
        problems += signature_placeholder_once(docs[n], 0)
        if n == RESPONSE:
            problems += strings_only_in_quotes(docs[n], REMOVED_TERMS)
        elif n != CHECKLIST:  # the checklist quotes the guidelines' own wording (RECORD 19.1: "misclassification bias")
            problems += strings_absent(docs[n], REMOVED_TERMS, ignore_case=True)
    for f in (ms, clean):
        problems += has_no_embedded_figures(f)
        problems += frontmatter_wordcount_matches(f)
        problems += abstract_within_limit(f) + declarations_present(f)
    if colored_run_count(clean):
        problems.append("clean manuscript contains colored runs")
    if not colored_run_count(ms):
        problems.append("colored manuscript has no colored runs")
    if highlight_count(ms) or highlight_count(clean):
        problems.append("highlight left in a manuscript")
    if _plain(ms) != _plain(clean):
        problems.append("colored and clean manuscripts differ in text")
    problems += no_sidecar_files(root)
    problems += tiff_ok(root / "Figures", FIG, MAP)
    problems += supplement_numbering(docs[SUPPLEMENT])
    problems += claims_in_docx(root)
    return problems


def main(root):
    """Print the verdict, then the open items; the exit status follows the problems alone."""
    problems = check(root)
    print("\n".join(problems) or "packet checks pass")
    for item in open_human_items(root):
        print("open item:", item)
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
