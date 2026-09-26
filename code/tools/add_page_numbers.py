"""Put a centred PAGE field in the footer of every section of a .docx (the house reference
document has no footer, and the journal asks for page numbers).

    python3 code/tools/add_page_numbers.py FILE.docx [...]
"""
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_page_numbers(path):
    doc = Document(path)
    for sec in doc.sections:
        p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        p.alignment = 1
        run = p.add_run()
        for kind, text in (("begin", None), ("instr", " PAGE "), ("end", None)):
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
    doc.save(path)


if __name__ == "__main__":
    for f in sys.argv[1:]:
        add_page_numbers(f)
