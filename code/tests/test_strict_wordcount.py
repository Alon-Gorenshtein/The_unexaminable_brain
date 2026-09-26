import docx
import pytest

import strict_wordcount as sw


def _build(path, with_tables_heading=True):
    d = docx.Document()
    d.add_heading("Introduction", level=1)
    d.add_paragraph("One two three four five.")
    d.add_heading("References", level=1)
    d.add_paragraph("1. Reference text that must not be counted.")
    if with_tables_heading:
        d.add_heading("Tables", level=1)
    d.add_paragraph("Table 1. Caption words here.")
    t = d.add_table(rows=2, cols=2)
    for i, row in enumerate([["alpha beta", "gamma"], ["delta", "epsilon zeta eta"]]):
        for j, cell in enumerate(row):
            t.cell(i, j).text = cell
    d.save(str(path))
    return path


def test_pandoc_shows_bare_section_headings(tmp_path):
    lines = sw._lines(_build(tmp_path / "doc.docx"))
    stripped = [ln.strip() for ln in lines]
    for name in ("Introduction", "References", "Tables"):
        assert name in stripped


def test_exact_count_counts_headings_and_cells_but_not_rules_or_references(tmp_path):
    # Introduction heading 1 + body 5 = 6
    # Tables heading 1 + caption 5 + cells (2 + 1 + 1 + 3 = 7) = 13
    # pandoc's dash-run table borders and the reference line add nothing
    assert sw.strict_count(_build(tmp_path / "doc.docx")) == 19


def test_missing_tables_heading_raises(tmp_path):
    path = _build(tmp_path / "no_tables.docx", with_tables_heading=False)
    with pytest.raises(ValueError, match="tables"):
        sw.strict_count(path)


def test_declarations_before_references_are_not_counted(tmp_path):
    d = docx.Document()
    d.add_heading("Introduction", level=1)
    d.add_paragraph("One two three four five.")
    d.add_heading("Author Contributions", level=1)
    d.add_paragraph("These words must not be counted at all.")
    d.add_heading("Statements and Declarations", level=1)
    d.add_paragraph("Neither must these words.")
    d.add_heading("References", level=1)
    d.add_paragraph("1. Reference text.")
    d.add_heading("Tables", level=1)
    d.add_paragraph("Table 1. Caption.")
    p = tmp_path / "decl.docx"
    d.save(str(p))
    assert sw.strict_count(str(p)) == 1 + 5 + 1 + 3   # "Introduction" heading, five words, "Tables" heading, three caption words


def test_acknowledgments_before_references_are_not_counted(tmp_path):
    d = docx.Document()
    d.add_heading("Introduction", level=1)
    d.add_paragraph("One two three four five.")
    d.add_heading("Acknowledgments", level=1)
    d.add_paragraph("Thanks to everyone who helped here.")
    d.add_heading("References", level=1)
    d.add_paragraph("1. Reference text.")
    d.add_heading("Tables", level=1)
    d.add_paragraph("Table 1. Caption.")
    p = tmp_path / "ack.docx"
    d.save(str(p))
    assert sw.strict_count(str(p)) == 1 + 5 + 1 + 3
