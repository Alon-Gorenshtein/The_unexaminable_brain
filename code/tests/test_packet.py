import os
import shutil
import time
import zipfile

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from PIL import Image

import check_packet as C
from add_page_numbers import add_page_numbers
from highlight_to_color import recolor
from set_docx_metadata import neutral


def _doc(path, text="Hello"):
    d = Document()
    d.add_paragraph(text)
    d.save(path)


def _rewrite(path, part, fn):
    """Apply fn to one XML part inside a .docx."""
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == part:
                data = fn(data.decode("utf8")).encode("utf8")
            zout.writestr(item, data)
    shutil.move(tmp, path)


def test_page_field_guard(tmp_path):
    a, b = tmp_path / "a.docx", tmp_path / "b.docx"
    _doc(a)
    _doc(b)
    add_page_numbers(b)
    assert C.has_page_field(a) and not C.has_page_field(b)


def test_double_space_guard(tmp_path):
    a, b, c = tmp_path / "a.docx", tmp_path / "b.docx", tmp_path / "c.docx"
    _doc(a)
    d = Document()
    d.add_paragraph("x").paragraph_format.line_spacing = 2
    d.save(b)
    d = Document()
    d.add_paragraph("x").paragraph_format.line_spacing = 2
    d.add_paragraph("y").paragraph_format.line_spacing = 1
    d.save(c)
    assert C.is_double_spaced(a) and not C.is_double_spaced(b)
    assert C.is_double_spaced(c)  # one paragraph overriding the double spacing trips it


def test_single_space_guard(tmp_path):
    a, b = tmp_path / "a.docx", tmp_path / "b.docx"
    d = Document()
    d.add_paragraph("x").paragraph_format.line_spacing = 2
    d.save(a)
    d = Document()
    d.add_paragraph("x").paragraph_format.line_spacing = 1
    d.save(b)
    assert C.is_single_spaced(a) and not C.is_single_spaced(b)


def test_figure_guard(tmp_path):
    Image.new("RGB", (4, 4)).save(tmp_path / "i.png")
    d = Document()
    d.add_picture(str(tmp_path / "i.png"))
    d.save(tmp_path / "f.docx")
    _doc(tmp_path / "g.docx")
    assert C.has_no_embedded_figures(tmp_path / "f.docx") and not C.has_no_embedded_figures(tmp_path / "g.docx")


def test_stale_guard(tmp_path):
    md, dx = tmp_path / "m.md", tmp_path / "m.docx"
    md.write_text("x")
    _doc(dx)
    os.utime(dx, (time.time() - 100, time.time() - 100))
    assert C.not_stale(dx, md)
    os.utime(dx, (time.time() + 100, time.time() + 100))
    assert not C.not_stale(dx, md)


def test_string_guards(tmp_path):
    p = tmp_path / "s.docx"
    _doc(p, "The default misclassifies severity.")
    assert C.strings_absent(p, ["misclassifies"]) and not C.strings_absent(p, ["no such word"])
    assert C.strings_absent(p, ["MISCLASSIF"], ignore_case=True)
    assert C.strings_present(p, ["not there"]) and not C.strings_present(p, ["default"])


def test_strings_only_in_quotes(tmp_path):
    good, bad = tmp_path / "g.docx", tmp_path / "b.docx"
    from docx.enum.style import WD_STYLE_TYPE
    d = Document()
    d.styles.add_style("Block Text", WD_STYLE_TYPE.PARAGRAPH)
    d.add_paragraph("We removed the word.")
    d.add_paragraph("The term misclassification implies a standard.", style="Block Text")
    d.save(good)
    d = Document()
    d.add_paragraph("The default misclassifies severity.")
    d.save(bad)
    assert not C.strings_only_in_quotes(good, ["misclassif"]) and C.strings_only_in_quotes(bad, ["misclassif"])


def test_signature_placeholder(tmp_path):
    p, q = tmp_path / "p.docx", tmp_path / "q.docx"
    _doc(p, "[Signature and typed name to be added by the corresponding author]")
    _doc(q, "Sincerely, The authors")
    assert not C.signature_placeholder_once(p, 1) and C.signature_placeholder_once(q, 1)
    assert C.signature_placeholder_once(p, 0) and not C.signature_placeholder_once(q, 0)


def _manuscript(path, stated, body_words=4):
    d = Document()
    d.add_paragraph(f"Word count: {stated}")
    d.add_paragraph("Introduction")
    d.add_paragraph(" ".join(["word"] * body_words))
    d.add_paragraph("References")
    d.add_paragraph("1. A reference.")
    d.add_paragraph("Tables")
    d.save(path)


def test_frontmatter_wordcount_guard(tmp_path):
    good, bad, none = tmp_path / "g.docx", tmp_path / "b.docx", tmp_path / "n.docx"
    _manuscript(good, 6)   # Introduction (1) + 4 words + Tables (1)
    _manuscript(bad, 99)
    _doc(none, "Introduction\n")
    assert not C.frontmatter_wordcount_matches(good)
    assert C.frontmatter_wordcount_matches(bad) and C.frontmatter_wordcount_matches(none)


def test_metadata_guard(tmp_path):
    p = tmp_path / "m.docx"
    _doc(p)
    d = Document(p)
    d.core_properties.author = "someone@example.org"
    d.save(p)
    assert C.no_forbidden_metadata(p)  # email address as author
    neutral(p)
    assert not C.no_forbidden_metadata(p)
    d = Document(p)
    d.core_properties.last_modified_by = "python-docx"
    d.save(p)
    assert C.no_forbidden_metadata(p)
    d = Document(p)
    d.core_properties.last_modified_by = "Some Tool"
    d.save(p)
    assert C.no_forbidden_metadata(p, [r"Some\b"]) and not C.no_forbidden_metadata(p, [r"Other\b"])
    neutral(p)
    assert not C.no_forbidden_metadata(p, [r"Some\b"])


def test_terms_file_and_metadata_terms(tmp_path):
    f = tmp_path / "terms.tsv"
    f.write_text("# comment\nprocess\tWIP\nmetadata\t\\bXY\\b\n")
    terms = C.load_terms(f)
    assert terms == {"process": ["WIP"], "metadata": [r"\bXY\b"]}
    p = tmp_path / "m.docx"
    _doc(p, "Draft WIP text")
    d = Document(p)
    d.core_properties.last_modified_by = "XY"
    d.save(p)
    assert C.no_forbidden_metadata(p, terms["metadata"]) and C.strings_absent(p, terms["process"])
    if C.TERMS_FILE.is_file():  # the real list is kept outside the published code
        real = C.load_terms()
        assert real["process"] and real["metadata"]


def test_sidecar_guard(tmp_path):
    (tmp_path / "Figures").mkdir()
    _doc(tmp_path / "a.docx")
    assert not C.no_sidecar_files(tmp_path)
    (tmp_path / "Figures" / "._Figure1.tif").write_bytes(b"x")
    assert C.no_sidecar_files(tmp_path)
    (tmp_path / "Figures" / "._Figure1.tif").unlink()
    with zipfile.ZipFile(tmp_path / "a.docx", "a") as z:
        z.writestr("word/._document.xml", b"x")
    assert C.no_sidecar_files(tmp_path)


def test_track_change_guard(tmp_path):
    p = tmp_path / "t.docx"
    _doc(p, "Hello")
    assert not C.no_track_changes(p)
    _rewrite(p, "word/document.xml",
             lambda x: x.replace("<w:r>", '<w:ins w:id="1" w:author="x" w:date="2026-01-01T00:00:00Z"><w:r>', 1)
             .replace("</w:r>", "</w:r></w:ins>", 1))
    assert C.no_track_changes(p)
    q = tmp_path / "c.docx"
    _doc(q, "Hello")
    _rewrite(q, "word/document.xml", lambda x: x.replace("</w:p>", '<w:commentRangeStart w:id="0"/></w:p>', 1))
    assert C.no_track_changes(q)


def test_tiff_guard(tmp_path):
    src, dst = tmp_path / "png", tmp_path / "Figures"
    src.mkdir()
    dst.mkdir()
    Image.new("RGBA", (20, 10), (200, 10, 10, 255)).save(src / "a.png")
    Image.new("RGB", (20, 10), (200, 10, 10)).save(dst / "F1.tif", dpi=(600, 600))
    assert not C.tiff_ok(dst, src, {"a": "F1"})
    Image.new("RGB", (20, 10), (200, 10, 10)).save(dst / "F1.tif", dpi=(300, 300))
    assert C.tiff_ok(dst, src, {"a": "F1"})                     # wrong dpi
    Image.new("RGBA", (20, 10), (200, 10, 10, 255)).save(dst / "F1.tif", dpi=(600, 600))
    assert C.tiff_ok(dst, src, {"a": "F1"})                     # alpha channel
    Image.new("RGB", (21, 10), (200, 10, 10)).save(dst / "F1.tif", dpi=(600, 600))
    assert C.tiff_ok(dst, src, {"a": "F1"})                     # different size
    Image.new("RGB", (20, 10), (0, 0, 0)).save(dst / "F1.tif", dpi=(600, 600))
    assert C.tiff_ok(dst, src, {"a": "F1"})                     # stale pixels
    assert C.tiff_ok(dst, src, {"a": "F1", "a2": "F2"})         # missing file


def test_supplement_numbering_guard(tmp_path):
    good, bad = tmp_path / "g.docx", tmp_path / "b.docx"
    d = Document()
    for i in range(1, 4):
        d.add_paragraph(f"eTable {i}. Caption {i}.")
        d.add_paragraph("body")
    d.add_paragraph("eFigure 1. A figure.")
    d.save(good)
    assert not C.supplement_numbering(good, n_tables=3, n_figures=1)
    d = Document()
    for i in (1, 3, 2):
        d.add_paragraph(f"eTable {i}. Caption {i}.")
    d.add_paragraph("eFigure 1. A figure.")
    d.save(bad)
    assert C.supplement_numbering(bad, n_tables=3, n_figures=1)   # out of order
    d = Document()
    for i in (1, 2):
        d.add_paragraph(f"eTable {i}. Caption {i}.")
    d.add_paragraph("eFigure 1. A figure.")
    d.save(bad)
    assert C.supplement_numbering(bad, n_tables=3, n_figures=1)   # missing eTable 3


def _supplement_docx(path, n_tables):
    d = Document()
    for i in range(1, n_tables + 1):
        d.add_paragraph(f"eTable {i}. Caption {i}.")
    for i in range(1, 4):
        d.add_paragraph(f"eFigure {i}. Figure {i}.")
    d.save(path)


def test_supplement_numbering_default_counts_every_generated_table(tmp_path):
    import importlib.util
    import inspect
    from config import PROJ
    spec = importlib.util.spec_from_file_location("supp_tables", PROJ / "code" / "74_build_supplement_tables.py")
    gen = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(gen)
    n = len(gen.TABLES)
    assert inspect.signature(C.supplement_numbering).parameters["n_tables"].default == n == 18
    good, short = tmp_path / "g.docx", tmp_path / "s.docx"
    _supplement_docx(good, n)
    _supplement_docx(short, n - 1)
    assert not C.supplement_numbering(good)
    assert C.supplement_numbering(short)      # an eTable missing from the packet supplement fails the default check


def test_docx_literal_and_text():
    assert C.docx_literal("**Eligible stays (n = 14,272)**") == "Eligible stays (n = 14,272)"
    assert C.docx_literal("stays)^b^ | 8683 | 736") == "stays)b | 8683 | 736"
    assert C.docx_literal("ICU^1,2^ with *P* < .001") == "ICU1,2 with P < .001"


@pytest.mark.requires_study_files("output/revision/revision_digest.json")
def test_claims_in_docx_guard(tmp_path):
    import json
    from pathlib import Path
    root = Path(C.ROOT)
    digest = json.loads((root / "output/revision/revision_digest.json").read_text())
    stem, key = next((s, k) for s, v in digest.items() if isinstance(v, dict)
                     for k, x in v.items() if isinstance(x, int) and not isinstance(x, bool) and x > 1000)
    from lib_fmt import fmt
    value = fmt(digest[stem][key], "int")
    claims = tmp_path / "claims.tsv"
    claims.write_text(f"manuscript/manuscript.md\t{stem}:{key}\tint\t**{value}** stays\n")
    d = Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Stays", str(value)
    d.add_paragraph(f"There were {value} stays.")
    d.save(tmp_path / C.CLEAN)
    assert not C.claims_in_docx(tmp_path, claims, scratch=tmp_path)
    _doc(tmp_path / C.CLEAN, "There were 1 stays.")
    assert C.claims_in_docx(tmp_path, claims, scratch=tmp_path)


def test_recolor_and_color_guards(tmp_path):
    p = tmp_path / "c.docx"
    d = Document()
    d.add_paragraph().add_run("changed").font.highlight_color = WD_COLOR_INDEX.YELLOW
    d.save(p)
    assert C.highlight_count(p) == 1 and C.colored_run_count(p) == 0
    assert recolor(p) == 1
    assert C.highlight_count(p) == 0 and C.colored_run_count(p) == 1


def test_recolor_reaches_runs_in_hyperlinks_and_tables(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p = tmp_path / "h.docx"
    d = Document()
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).paragraphs[0].add_run("cell").font.highlight_color = WD_COLOR_INDEX.YELLOW
    para = d.add_paragraph()
    run = para.add_run("link text")
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "x")
    para._p.append(link)
    link.append(run._r)  # the run now sits inside the hyperlink, not directly in the paragraph
    d.save(p)
    assert C.highlight_count(p) == 2 and "<w:hyperlink" in C._xml(p, "word/document.xml")
    assert recolor(p) == 2 and C.highlight_count(p) == 0 and C.colored_run_count(p) == 2


@pytest.mark.requires_study_files("manuscript/reporting_checklist.md")
def test_checklist_is_registered():
    assert C.CHECKLIST == "5_Reporting_Checklist_STROBE_RECORD.docx"
    assert C.CHECKLIST in C.DOCS and len(C.DOCS) == 6
    assert C.SOURCES[C.CHECKLIST] == "reporting_checklist.md"
    assert (C.ROOT / "manuscript" / C.SOURCES[C.CHECKLIST]).is_file()


def test_checklist_title_guard(tmp_path):
    title = tmp_path / "title.txt"
    title.write_text("A Title\n")
    bad, good = tmp_path / "bad.docx", tmp_path / "good.docx"
    _doc(bad, "STROBE and RECORD Reporting Checklist")
    _doc(good, "A Title")
    assert C.title_matches(bad, title) and not C.title_matches(good, title)


def _packet(root, texts, spacing):
    """A packet of the six documents, each with the real title, a page field and the given spacing."""
    title = (C.ROOT / "manuscript" / "title.txt").read_text().strip()
    for n in C.DOCS:
        d = Document()
        for t in (title, texts.get(n, "Body.")):
            d.add_paragraph(t).paragraph_format.line_spacing = spacing.get(n, 1)
        d.save(root / n)
        add_page_numbers(root / n)


@pytest.mark.requires_study_files("manuscript/title.txt")
def test_check_requires_the_checklist(tmp_path):
    _packet(tmp_path, {}, {})
    (tmp_path / C.CHECKLIST).unlink()
    assert C.check(tmp_path) == [f"missing {C.CHECKLIST}"]


@pytest.mark.requires_study_files("manuscript/title.txt")
def test_check_spacing_and_removed_terms_for_the_checklist(tmp_path, monkeypatch):
    for f in ("tiff_ok", "supplement_numbering", "claims_in_docx", "frontmatter_wordcount_matches"):
        monkeypatch.setattr(C, f, lambda *a, **k: [])
    quote = "Include discussion of misclassification bias."
    _packet(tmp_path, {C.CHECKLIST: quote, C.COVER: quote}, {C.CHECKLIST: 2})
    about = lambda name: [p for p in C.check(tmp_path) if p.startswith(name)]
    # the checklist quotes the guideline, so the removed-terms scan skips it; a letter is still caught
    assert not [p for p in about(C.CHECKLIST) if "misclassif" in p]
    assert [p for p in about(C.COVER) if "misclassif" in p]
    # the checklist must be single-spaced, like the letters
    assert [p for p in about(C.CHECKLIST) if "not single-spaced" in p]


def _abstract_doc(path, n_words):
    d = Document()
    d.add_paragraph("Abstract")
    d.add_paragraph(" ".join(["word"] * n_words))
    d.add_paragraph("Keywords: a; b")
    d.save(path)


@pytest.mark.parametrize("n,ok", [(249, False), (250, True), (300, True), (301, False), (393, False)])
def test_abstract_limit_guard(tmp_path, n, ok):
    p = tmp_path / "a.docx"
    _abstract_doc(p, n)
    assert (C.abstract_within_limit(p) == []) is ok


def test_abstract_limit_message_names_count_and_range(tmp_path):
    p = tmp_path / "a.docx"
    _abstract_doc(p, 393)
    assert C.abstract_words(p) == 393
    assert C.abstract_within_limit(p) == ["a.docx: abstract has 393 words; the journal requires 250-300"]


def test_abstract_limit_needs_heading_and_keywords_line(tmp_path):
    no_keywords, no_heading, wrong_order = (tmp_path / n for n in ("k.docx", "h.docx", "o.docx"))
    d = Document()
    d.add_paragraph("Abstract")
    d.add_paragraph(" ".join(["word"] * 280))
    d.save(no_keywords)
    d = Document()
    d.add_paragraph(" ".join(["word"] * 280))
    d.add_paragraph("Keywords: a; b")
    d.save(no_heading)
    d = Document()
    d.add_paragraph("Keywords: a; b")
    d.add_paragraph("Abstract")
    d.add_paragraph(" ".join(["word"] * 280))
    d.save(wrong_order)
    for p in (no_keywords, no_heading, wrong_order):
        assert C.abstract_words(p) == -1
        assert C.abstract_within_limit(p) == [f"{p.name}: abstract or keywords line not found"]


def test_abstract_count_ignores_a_keywords_line_on_the_title_page(tmp_path):
    p = tmp_path / "t.docx"
    d = Document()
    d.add_paragraph("Keywords: title page copy")
    d.add_paragraph("Abstract")
    d.add_paragraph(" ".join(["word"] * 280))
    d.add_paragraph("Keywords: a; b")
    d.save(p)
    assert C.abstract_words(p) == 280 and C.abstract_within_limit(p) == []


def test_abstract_count_includes_section_labels_and_stops_at_keywords(tmp_path):
    p = tmp_path / "l.docx"
    d = Document()
    d.add_paragraph("Title words that are not part of the abstract")
    d.add_paragraph("Abstract")
    d.add_paragraph("Background: one two three")
    d.add_paragraph("Conclusion: four five")
    d.add_paragraph("Keywords: a; b; c")
    d.add_paragraph("Introduction")
    d.add_paragraph("Body words after the keywords")
    d.save(p)
    assert C.abstract_words(p) == 7


ALL_DECLARATIONS = ["Acknowledgments", "Author Contributions", "Statements and Declarations",
                    "Ethical considerations", "Consent to participate", "Consent for publication",
                    "Declaration of conflicting interest", "Funding statement", "Data availability"]


def test_declarations_guard_reports_each_missing_heading(tmp_path):
    p = tmp_path / "d.docx"
    d = Document()
    d.add_paragraph("Conclusion")
    d.add_paragraph("Author Contributions")
    d.add_paragraph("Statements and Declarations")
    d.add_paragraph("Ethical considerations")
    d.save(p)
    problems = C.declarations_present(p)
    assert any("Consent to participate" in x for x in problems) and any("Data availability" in x for x in problems)
    assert not any("Author Contributions" in x for x in problems)
    assert len(problems) == 6


def test_declarations_guard_passes_with_all_nine_headings(tmp_path):
    p = tmp_path / "d.docx"
    d = Document()
    for h in ALL_DECLARATIONS:
        d.add_paragraph(h)
        d.add_paragraph("Text under the heading.")
    d.save(p)
    assert C.declarations_present(p) == []


def test_declarations_guard_wants_headings_as_their_own_lines(tmp_path):
    p = tmp_path / "d.docx"
    d = Document()
    d.add_paragraph("The Funding statement and Data availability are given elsewhere.")
    d.save(p)
    assert len(C.declarations_present(p)) == 9


def test_declarations_guard_reports_a_missing_acknowledgments_heading(tmp_path):
    p = tmp_path / "d.docx"
    d = Document()
    for h in ALL_DECLARATIONS[1:]:
        d.add_paragraph(h)
    d.save(p)
    assert C.declarations_present(p) == ["d.docx: missing heading 'Acknowledgments'"]


def test_open_human_items_lists_the_marker(tmp_path):
    d = Document()
    d.add_paragraph("[CRediT roles to be added by the authors]")
    d.save(tmp_path / "3_Manuscript_COLORED_CHANGES.docx")
    assert C.open_human_items(tmp_path) == ["3_Manuscript_COLORED_CHANGES.docx: [CRediT roles to be added by the authors]"]


def test_open_human_items_is_empty_without_the_marker_or_files(tmp_path):
    assert C.open_human_items(tmp_path) == []
    _doc(tmp_path / "3_Manuscript_COLORED_CHANGES.docx", "Author Contributions")
    assert C.open_human_items(tmp_path) == []


@pytest.mark.requires_study_files("manuscript/title.txt")
def test_check_runs_the_abstract_and_declaration_guards_on_both_manuscripts(tmp_path, monkeypatch):
    for f in ("tiff_ok", "supplement_numbering", "claims_in_docx", "frontmatter_wordcount_matches"):
        monkeypatch.setattr(C, f, lambda *a, **k: [])
    _packet(tmp_path, {}, {})
    about = lambda name, word: [p for p in C.check(tmp_path) if p.startswith(name) and word in p]
    for n in (C.COLORED, C.CLEAN):
        assert about(n, "abstract or keywords line not found") and about(n, "missing heading 'Data availability'")
    for n in (C.COVER, C.RESPONSE, C.SUPPLEMENT, C.CHECKLIST):
        assert not about(n, "abstract") and not about(n, "missing heading")


def _marked_packet(root):
    d = Document()
    d.add_paragraph("[CRediT roles to be added by the authors]")
    d.save(root / C.COLORED)


def test_main_prints_open_items_after_a_pass_and_exits_zero(tmp_path, monkeypatch, capsys):
    _marked_packet(tmp_path)
    monkeypatch.setattr(C, "check", lambda root: [])
    assert C.main(tmp_path) == 0
    assert capsys.readouterr().out.splitlines() == [
        "packet checks pass", f"open item: {C.COLORED}: [CRediT roles to be added by the authors]"]


def test_main_exits_one_on_a_problem_and_never_claims_a_pass(tmp_path, monkeypatch, capsys):
    _marked_packet(tmp_path)
    monkeypatch.setattr(C, "check", lambda root: ["missing 5_x.docx"])
    assert C.main(tmp_path) == 1
    out = capsys.readouterr().out
    assert "packet checks pass" not in out and "missing 5_x.docx" in out


def test_main_open_items_do_not_change_the_exit_status(tmp_path, monkeypatch):
    monkeypatch.setattr(C, "check", lambda root: [])
    assert C.main(tmp_path) == 0
    _marked_packet(tmp_path)
    assert C.open_human_items(tmp_path) and C.main(tmp_path) == 0
