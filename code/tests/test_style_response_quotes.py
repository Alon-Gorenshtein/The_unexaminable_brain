from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import RGBColor

from style_response_quotes import BLUE, style


def _letter(path):
    d = Document()
    d.styles.add_style("Block Text", WD_STYLE_TYPE.PARAGRAPH)
    d.add_heading("Comment 1.2.", level=3)
    q = d.add_paragraph(style="Block Text")
    q.add_run("Reviewer comment.").font.color.rgb = BLUE  # as the house script leaves it
    d.add_paragraph().add_run("Response.").bold = True
    d.add_paragraph().add_run("Original:").bold = True
    d.add_paragraph(style="Block Text").add_run("Submitted text.").font.color.rgb = BLUE
    d.add_paragraph().add_run("Revised:").bold = True
    d.add_paragraph(style="Block Text").add_run("Revised text.").italic = True
    tb = d.add_table(rows=1, cols=2)
    tb.cell(0, 0).paragraphs[0].add_run("Revised cell")
    d.add_paragraph().add_run("Original:").bold = True
    tb = d.add_table(rows=1, cols=1)
    tb.cell(0, 0).paragraphs[0].add_run("Submitted cell")
    d.add_paragraph("Ordinary prose ends the quotation.")
    tb = d.add_table(rows=1, cols=1)
    tb.cell(0, 0).paragraphs[0].add_run("Letter cell")
    d.add_heading("Comment 1.3.", level=3)
    d.add_paragraph(style="Block Text").add_run("Next comment.").font.color.rgb = BLUE
    d.save(path)


def test_quotes_styled_by_label(tmp_path):
    p = tmp_path / "r.docx"
    _letter(p)
    assert style(p) == (1, 3, 2)
    doc = Document(p)
    runs = {r.text: r for para in doc.paragraphs for r in para.runs}
    runs.update({r.text: r for t in doc.tables for c in t._cells for r in c.paragraphs[0].runs})
    assert runs["Revised cell"].font.color.rgb == BLUE and runs["Submitted cell"].italic
    assert runs["Letter cell"].font.color.rgb is None and not runs["Letter cell"].italic
    for t in ("Reviewer comment.", "Submitted text.", "Next comment."):
        assert runs[t].font.color.rgb is None and runs[t].italic
    assert runs["Revised text."].font.color.rgb == RGBColor(0x1F, 0x4E, 0x79) and not runs["Revised text."].italic
